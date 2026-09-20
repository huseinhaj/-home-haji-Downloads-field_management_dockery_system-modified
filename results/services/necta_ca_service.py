"""
necta_ca_service.py — NECTA Continuous Assessment Form (Form IV) automation.

Baraza la NECTA huomba Excel ya CONTINUOUS ASSESSMENT FORM FOR SECONDARY
SCHOOLS: kila somo form yake, wanafunzi waliosajiliwa ndiyo kwanza, kisha
columns za FIRST TERM (TEST-ONE, MID TERM, TEST-TWO, TERMINAL), SECOND
TERM (TEST-ONE, MID-TERM, TEST-TWO, ANNUAL), PROJECT na PRACTICAL.

Automation hapa:
  1. Roster ya Form 4 (is_active, mwaka wa sasa) inajaza S/N + majina.
  2. Mwalimu anachagua somo + exam ya kila column (auto-map kwa exam_type)
     + range ya alama (mf. 45-100) — range hii ni KIKOMO (floor/ceiling).
  3. Alama zinakadiriwa kwa BANDING YA MWANAFUNZI (performance-aware):
     kila mwanafunzi ana alama yake ya CHINI (student_min) kutoka kwenye
     mitihani iliyochaguliwa; floor = max(mark_min, student_min).
         estimate(b) = floor + (b - student_min) * (mark_max - floor)
                       / (100 - student_min)
     - Mwanafunzi dhaifu (alama ya chini 40, range 45-100) huanzia 45 na
       kupanda kadiri utendaji wake — hawezi kutiwa chini ya 45.
     - Mwanafunzi hodari (alama ya chini 70, range 45-100) huanzia 70
       kwenda juu — hashushishwi hadi 45.
     Msingi wa kila column ni alama ya exam ile column; kama column ile
     haina exam/data, inatumika wastani wa combined results za mwanafunzi
     kutoka columns zote zilizopo.
  4. PROJECT na PRACTICAL zinabaki WAZI (mwalimu anajaza mkononi).
  5. Excel inatengenezwa kwa openpyxl kwa muundo wa NECTA.
"""
from __future__ import annotations

import math

import openpyxl
from django.http import HttpResponse
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ── Column definitions (NECTA C.A. layout) ─────────────────────────────────
# Kila column ya muhula ina exam_types ambazo zinafaa kuilisha — auto-map
# inachagua exam ya kwanza isiyotumika yenye type hiyo.
TERM_COLUMNS = [
    {'key': 'ft_test_one', 'term': 'FIRST TERM',  'label': 'TEST-ONE',  'types': ['TEST', 'MONTHLY', 'QUIZ']},
    {'key': 'ft_mid_term', 'term': 'FIRST TERM',  'label': 'MID TERM',  'types': ['MIDTERM']},
    {'key': 'ft_test_two', 'term': 'FIRST TERM',  'label': 'TEST-TWO',  'types': ['TEST', 'MONTHLY', 'QUIZ']},
    {'key': 'ft_terminal', 'term': 'FIRST TERM',  'label': 'TERMINAL',  'types': ['TERMINAL']},
    {'key': 'st_test_one', 'term': 'SECOND TERM', 'label': 'TEST-ONE',  'types': ['TEST', 'MONTHLY', 'QUIZ']},
    {'key': 'st_mid_term', 'term': 'SECOND TERM', 'label': 'MID-TERM',  'types': ['MIDTERM']},
    {'key': 'st_test_two', 'term': 'SECOND TERM', 'label': 'TEST-TWO',  'types': ['TEST', 'MONTHLY', 'QUIZ']},
    {'key': 'st_annual',   'term': 'SECOND TERM', 'label': 'ANNUAL',    'types': ['ANNUAL']},
]
EXTRA_COLUMNS = [
    {'key': 'project',   'term': '', 'label': 'PROJECT',   'types': []},   # wazi — manual
    {'key': 'practical', 'term': '', 'label': 'PRACTICAL', 'types': []},   # wazi — manual
]
ALL_COLUMNS = TERM_COLUMNS + EXTRA_COLUMNS
COLUMN_KEYS = [c['key'] for c in ALL_COLUMNS]


def auto_map_exams(form4_exams):
    """Chagua exam ya kila column kwa exam_type (ya kwanza isiyotumika).

    form4_exams: queryset/dict-like ya Exam za Form 4 (school-scoped).
    Returns {column_key: exam_id or None}.
    """
    available = list(form4_exams)
    mapping = {}
    used = set()
    for col in TERM_COLUMNS:
        pick = next(
            (e for e in available
             if e.id not in used and (e.exam_type in col['types'])),
            None,
        )
        if pick is not None:
            used.add(pick.id)
        mapping[col['key']] = pick.id if pick else None
    # EXTRA columns (PROJECT/PRACTICAL) hazina exams — zinaachiwa wazi.
    for col in EXTRA_COLUMNS:
        mapping[col['key']] = None
    return mapping


def estimate_mark(base_score, mark_min, mark_max, student_min=None):
    """Banding ya mwanafunzi ndani ya range ya mwalimu (floor/ceiling).

    student_min = alama ya CHINI ya mwanafunzi kwenye mitihani
    iliyochaguliwa (profile yake ya utendaji):
      - Dhaifu (student_min < mark_min, mf. 40 na range 45-100) → huanzia
        mark_min (45) na kupanda kadiri utendaji wake.
      - Hodari (student_min 70 na range 45-100) → huanzia 70 kwenda juu;
        hashushishwi hadi 45.

    estimate(b) = floor + (b - student_min) * (mark_max - floor)
                  / (100 - student_min),   floor = max(mark_min, student_min)

    student_min=None → fallback ya proportional: min + (b/100)*(max-min).
    base_score None → None. Matokeo yanafungwa kwenye [mark_min, mark_max].
    """
    if base_score is None:
        return None
    mark_min = int(mark_min)
    mark_max = int(mark_max)
    base = float(base_score)

    if student_min is None:
        mark = mark_min + (base / 100.0) * (mark_max - mark_min)
    else:
        smin = min(max(float(student_min), 0.0), 100.0)
        floor = max(mark_min, smin)
        span_real = 100.0 - smin
        if span_real <= 0:
            mark = mark_max          # alama zote za mwanafunzi ni 100
        elif base <= smin:
            mark = floor             # alama ya chini → floor
        else:
            mark = floor + (base - smin) * (mark_max - floor) / span_real

    mark = math.floor(mark + 0.5)
    return max(mark_min, min(mark_max, mark))


def _full_name(fs):
    return ' '.join(p for p in [fs.first_name, fs.middle_name, fs.last_name] if p)


def _match_students(roster_rows):
    """Funga kila FormStudent row na Student record yake (alama ziko kwenye
    Student). Inalinganisha kwa (first, middle, last) — kaka wenye
    first+last moja hutofautishwa na middle name; kama hakuna match kamili,
    (first, last) pekee inatumika tu ikiwa ni unique."""
    from ..models import Student

    firsts = {fs.first_name for fs in roster_rows}
    lasts = {fs.last_name for fs in roster_rows}
    by_full = {}
    by_fl = {}
    for s in Student.objects.filter(first_name__in=firsts, last_name__in=lasts):
        by_full.setdefault(
            (s.first_name.lower(), (s.middle_name or '').lower(), s.last_name.lower()), s,
        )
        by_fl.setdefault((s.first_name.lower(), s.last_name.lower()), []).append(s)

    matched = {}
    for fs in roster_rows:
        key_full = (fs.first_name.lower(), (fs.middle_name or '').lower(), fs.last_name.lower())
        student = by_full.get(key_full)
        if student is None:
            candidates = by_fl.get((fs.first_name.lower(), fs.last_name.lower()), [])
            if len(candidates) == 1:
                student = candidates[0]
        matched[fs.id] = student
    return matched


def build_ca_preview(school, subject, exam_map, mark_min, mark_max, year=None,
                     form_num=4):
    """Tengeneza grid ya C.A. form: roster ya Form (form_num) + alama zilizokadiriwa.

    Returns dict:
      columns   — metadata ya columns (kwa header ya Excel/preview)
      exam_map  — {column_key: exam_id or None}
      rows      — [{'sn', 'name', 'admission_no', 'marks': {key: int|None}}]
      stats     — {filled, blank, students}
    """
    from ..models import Exam, ExamResult, FormStudent

    mark_min = int(mark_min)
    mark_max = int(mark_max)

    year = year or school.current_academic_year
    roster = FormStudent.objects.filter(
        school=school, form=form_num, is_active=True,
    )
    if year:
        roster = roster.filter(academic_year=year)
    if not roster.exists():
        # Fallback: roster yoyote ya form hiyo iliyopo (mwaka hazipangwi)
        roster = FormStudent.objects.filter(school=school, form=form_num, is_active=True)
    roster = roster.order_by('last_name', 'first_name', 'id')

    exam_ids = [eid for eid in exam_map.values() if eid]
    exam_by_id = {}
    if exam_ids:
        for e in Exam.objects.filter(school=school, id__in=exam_ids):
            exam_by_id[e.id] = e

    matched = _match_students(list(roster))
    student_ids = [s.id for s in matched.values() if s is not None]

    # {exam_id: {student_id: score}} — absent (X) inachukuliwa kama hakuna
    scores_by_exam = {}
    if student_ids:
        for er in (ExamResult.objects
                   .filter(exam_id__in=exam_ids, student_id__in=student_ids,
                           subject=subject, is_absent=False, score__isnull=False)
                   .values_list('exam_id', 'student_id', 'score')):
            scores_by_exam.setdefault(er[0], {})[er[1]] = er[2]

    rows = []
    filled = 0
    blank = 0
    for i, fs in enumerate(roster, start=1):
        student = matched.get(fs.id)
        sid = student.id if student else None

        column_scores = {}
        all_scores = []          # profile ya utendaji: alama halisi zote
        for col in TERM_COLUMNS:
            eid = exam_map.get(col['key'])
            if eid and sid:
                score = scores_by_exam.get(eid, {}).get(sid)
                if score is not None:
                    column_scores[col['key']] = float(score)
                    all_scores.append(float(score))

        combined_avg = (sum(all_scores) / len(all_scores)) if all_scores else None
        # Alama ya chini ya mwanafunzi kwenye mitihani iliyochaguliwa —
        # ndiyo anapoanzia makadirio (angalia estimate_mark).
        student_min = min(all_scores) if all_scores else None

        marks = {}
        for col in ALL_COLUMNS:
            if col['key'] in ('project', 'practical'):
                marks[col['key']] = None     # wazi — kujaza mkononi
                continue
            base = column_scores.get(col['key'], combined_avg)
            marks[col['key']] = estimate_mark(base, mark_min, mark_max, student_min)

        if any(v is not None for v in marks.values()):
            filled += 1
        else:
            blank += 1

        rows.append({
            'sn': i,
            'name': _full_name(fs),
            'admission_no': fs.admission_no or '',
            'marks': marks,
        })

    return {
        'columns': ALL_COLUMNS,
        'exam_map': {k: exam_map.get(k) for k in COLUMN_KEYS},
        'rows': rows,
        'stats': {'filled': filled, 'blank': blank, 'students': len(rows)},
    }


# ── Excel generation (NECTA layout) ────────────────────────────────────────
_THIN = Side(style='thin', color='FF000000')
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill('solid', fgColor='FFD9E1F2')
_ROMAN = {1: 'I', 2: 'II', 3: 'III', 4: 'IV', 5: 'V', 6: 'VI'}
_TITLE_FONT = Font(bold=True, size=14)
_SUBTITLE_FONT = Font(bold=True, size=12)
_INFO_FONT = Font(bold=True, size=11)


def generate_ca_excel(school, subject_name, year, rows, exam_map,
                      center_no='', phone='', form_num=4):
    """Tengeneza Excel ya NECTA Continuous Assessment Form.

    Layout (sheet 'C.A. FORM'):
      Row 1: THE NATIONAL EXAMINATIONS COUNCIL OF TANZANIA
      Row 2: CONTINUOUS ASSESSMENT FORM FOR SECONDARY SCHOOLS
      Row 3: FORM | PHONE NO | CENTER NUMBER | YEAR | SUBJECT — kila moja
             cell yake; phone/center ni TEXT ('@') ili namba ndefu
             zisibandikwe kama 2.6E+11 na 0 ya mwanzo isikatike
      Rows 4-5: headers (merged FIRST TERM / SECOND TERM groups)
      Rows 6+: S/N, NAME OF STUDENT, alama (PROJECT/PRACTICAL wazi)
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'C.A. FORM'

    n_cols = 12   # S/N + NAME + 4 + 4 + PROJECT + PRACTICAL

    # ── Title rows ──
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    c = ws.cell(row=1, column=1, value='THE NATIONAL EXAMINATIONS COUNCIL OF TANZANIA')
    c.font = _TITLE_FONT
    c.alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    c = ws.cell(row=2, column=1, value='CONTINUOUS ASSESSMENT FORM FOR SECONDARY SCHOOLS')
    c.font = _SUBTITLE_FONT
    c.alignment = Alignment(horizontal='center', vertical='center')

    # ── Row 3: FORM | PHONE NO | CENTER NUMBER | YEAR | SUBJECT ──
    # Phone na Center Number ni cells za TEXT (number_format '@') — hii
    # inazuia Excel/WPS kuandika namba ndefu kama scientific notation
    # (mf. 255712345678 → 2.6E+11) au kukata 0 ya mwanzo (0712… → 712…).
    try:
        rom = _ROMAN.get(int(form_num), str(form_num))
    except (TypeError, ValueError):
        rom = str(form_num or 4)

    def _field_label(col, label, merge_to=None):
        if merge_to and merge_to > col:
            ws.merge_cells(start_row=3, start_column=col,
                           end_row=3, end_column=merge_to)
        c = ws.cell(row=3, column=col, value=label)
        c.font = _INFO_FONT
        c.alignment = Alignment(horizontal='left', vertical='center')

    def _field_value(col, value, text=False):
        c = ws.cell(row=3, column=col, value=value)
        c.font = _INFO_FONT
        c.alignment = Alignment(horizontal='left', vertical='center')
        if text:
            c.number_format = '@'

    _field_label(1, f'FORM: {rom}', merge_to=2)
    _field_label(3, 'PHONE NO:')
    _field_value(4, phone or '……………', text=True)
    _field_label(5, 'CENTER NUMBER:', merge_to=6)
    _field_value(7, center_no or '……………', text=True)
    _field_label(8, 'YEAR:')
    _field_value(9, year)
    _field_label(10, 'SUBJECT:')
    _field_label(11, (subject_name or '').upper(), merge_to=n_cols)

    for col in range(1, n_cols + 1):
        ws.cell(row=3, column=col).border = _BORDER

    # ── Header rows 4-5 ──
    ws.merge_cells('A4:A5')
    ws['A4'] = 'S/N'
    ws.merge_cells('B4:B5')
    ws['B4'] = 'NAME OF STUDENT'

    first_term_cols = [c for c in TERM_COLUMNS if c['term'] == 'FIRST TERM']
    second_term_cols = [c for c in TERM_COLUMNS if c['term'] == 'SECOND TERM']

    def _write_group(start_col, term_label, cols):
        end_col = start_col + len(cols) - 1
        ws.merge_cells(start_row=4, start_column=start_col, end_row=4, end_column=end_col)
        ws.cell(row=4, column=start_col, value=term_label)
        for offset, col in enumerate(cols):
            ws.cell(row=5, column=start_col + offset, value=col['label'])
        return end_col + 1

    next_col = _write_group(3, 'FIRST TERM', first_term_cols)
    next_col = _write_group(next_col, 'SECOND TERM', second_term_cols)

    # PROJECT / PRACTICAL — wazi, mwalimu anajaza mkononi
    for label in ('PROJECT', 'PRACTICAL'):
        ws.merge_cells(start_row=4, start_column=next_col, end_row=5, end_column=next_col)
        ws.cell(row=4, column=next_col, value=label)
        next_col += 1

    for row in (4, 5):
        for col in range(1, n_cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = Font(bold=True, size=10)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = _BORDER
            cell.fill = _HEADER_FILL

    # ── Data rows ──
    for r_idx, row in enumerate(rows, start=6):
        ws.cell(row=r_idx, column=1, value=row['sn'])
        ws.cell(row=r_idx, column=2, value=(row['name'] or '').upper())
        col_idx = 3
        for key in COLUMN_KEYS:
            ws.cell(row=r_idx, column=col_idx, value=row['marks'].get(key))
            col_idx += 1
        for col in range(1, n_cols + 1):
            cell = ws.cell(row=r_idx, column=col)
            cell.border = _BORDER
            if col != 2:
                cell.alignment = Alignment(horizontal='center')

    # ── Column widths + print setup ──
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 34
    for col in range(3, n_cols + 1):
        ws.column_dimensions[get_column_letter(col)].width = 12
    ws.freeze_panes = 'A6'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    safe_subject = (subject_name or 'subject').replace(' ', '_').replace('/', '-')
    response['Content-Disposition'] = (
        f'attachment; filename="NECTA_CA_Form{form_num}_{safe_subject}_{year}.xlsx"'
    )
    wb.save(response)
    return response
