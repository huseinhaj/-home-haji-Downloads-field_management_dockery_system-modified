"""
Excel za Halmashauri kwa joint exam ya wilaya — NAKALA ya format za mafaili
ya Halmashauri (templates ziko results/xlsx_templates/, zimetolewa kwenye
mafaili halisi ya Kyerwa DC: vichwa, merges, Arial, rangi za theme, upana
wa safu, number formats):

  joint_division.xlsx  → "FORM ONE DIVISION PERFORMANCE ANALYSIS"
      DIVISION  — madaraja kwa kila shule + NAFASI KIWILAYA (RANK)
      GRADE     — gredi ya wastani wa mwanafunzi A–F (F/M/T)
  joint_subjects.xlsx  → "SUBJECT PERFORMANCE ANALYSIS"
      LIST OF SUBJECTS — ranking ya masomo (inasoma TOTAL ya kila sheet)
      <somo>           — sheet moja kwa kila somo (H'MAADILI, PHYS, KISW, ...)

Mfumo unaandika IDADI tu (WAV/WAS kwa kila daraja/gredi); jumla, %, GPA na
nafasi ni formulas za Excel zile zile za mafaili ya Halmashauri — faili
likifunguliwa Excel inahesabu, na afisa anaweza kuhariri kama kawaida.

Tofauti moja ya makusudi: mstari wa TOTAL wa DIVISION kwenye faili la
Halmashauri ulikuwa na makosa mawili (I–III ilijumlisha Div IV; IV–0 JML
ilijumlisha safu ya WAV) — hapa zimesahihishwa.
"""
from __future__ import annotations

import copy
import io
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as L

from .joint_analysis import GRADES, analyse_joint_exam

TEMPLATE_DIR = Path(__file__).resolve().parent.parent / 'xlsx_templates'

FORM_SW = {1: 'KIDATO CHA KWANZA', 2: 'KIDATO CHA PILI', 3: 'KIDATO CHA TATU',
           4: 'KIDATO CHA NNE', 5: 'KIDATO CHA TANO', 6: 'KIDATO CHA SITA'}
FORM_EN = {1: 'FORM ONE', 2: 'FORM TWO', 3: 'FORM THREE', 4: 'FORM FOUR',
           5: 'FORM FIVE', 6: 'FORM SIX'}
MONTH_SW = ['JANUARI', 'FEBRUARI', 'MACHI', 'APRILI', 'MEI', 'JUNI', 'JULAI',
            'AGOSTI', 'SEPTEMBA', 'OKTOBA', 'NOVEMBA', 'DESEMBA']
MONTH_EN = ['JANUARY', 'FEBRUARY', 'MARCH', 'APRIL', 'MAY', 'JUNE', 'JULY',
            'AUGUST', 'SEPTEMBER', 'OCTOBER', 'NOVEMBER', 'DECEMBER']
OWNERSHIP_LABEL = {'GOV': 'GOV', 'PRIVATE': 'PRIVATE', None: 'ALL'}

# Jina la sheet + jina kamili la somo kama kwenye faili la Halmashauri.
# Key = jina la somo bila alama/nafasi, herufi ndogo (majina ya DB yanatofautiana).
SUBJECT_SHEETS = [
    ("H'MAADILI", 'HISTORIA YA TANZANIA & MAADILI',
     ['historiayatanzanianamaadili', 'historiayatanzaniamaadili', 'hisroriayatanzanianamahadili']),
    ('PHYS', 'PHYSICS', ['physics']),
    ('EDK', 'ELIMU YA DINI YA KIISLAMU (EDK)',
     ['edk', 'elimuyadini', 'elimuyadiniyakiislam', 'elimuyadiniyakiislamu', 'islamicknowledge', 'ire']),
    ('HIST', 'HISTORY', ['history']),
    ('BIOS', 'BIOLOGY', ['biology']),
    ('CHEM', 'CHEMISTRY', ['chemistry']),
    ('MATHS', 'MATHEMATICS', ['mathematics', 'basicmathematics', 'basicappliedmathematics', 'hisabati', 'hesabu']),
    ('ENGLISH', 'ENGLISH', ['english', 'englishlanguage']),
    ('KISW', 'KISWAHILI', ['kiswahili']),
    ('GEO', 'GEOGRAPHY', ['geography']),
    ("BUS'STUDIES", 'BUSSINESS STUDIES', ['businessstudies', 'bussinessstudies']),
    ("B'KNOWLEDGE", 'BIBLE KNOWLEDGE', ['bibleknowledge', 'cre']),
    ("B'KEEPING", 'BOOK KEEPING', ['bookkeeping', 'bkeeping']),
    ('AGR', 'AGRICULTURE', ['agriculture']),
    ("LIT'ENGLISH", 'LITERATURE IN ENGLISH', ['literatureinenglish']),
    ("AD'MATHS", 'ADDITIONAL MATHEMATICS', ['additionalmathematics', 'furthermathematics']),
    ('FRENCH', 'FRENCH', ['french']),
    ("COMP'SCIENCE", 'COMPUTER SCIENCE', ['computerscience', 'computerstudies']),
    ("FAS'KISWAHILI", 'FASIHI YA KISWAHILI', ['fasihiyakiswahili']),
    ('MUSIC', 'MUSIC', ['music', 'muziki']),
    ('CHINESE', 'CHINESE', ['chinese']),
    ('ARABIC', 'ARABIC', ['arabic']),
    ("T'ARTS", 'THEATRE ARTS', ['theatrearts']),
    ("F'ARTS", 'FINE ARTS', ['finearts', 'fineart']),
    ("TG'CONSTR", 'TEXTILE & GARMENT CONSTRUCTION', ['textilegarmentconstruction', 'textileandgarmentconstruction']),
    ("SP'STUDIES", 'SPORT STUDIES', ['sportstudies', 'sportsstudies']),
    ("F&H'NUTR", 'FOOD & HUMAN NUTRITION', ['foodhumannutrition', 'foodandhumannutrition']),
]
_SUBJECT_LOOKUP = {k: (i, sheet, full) for i, (sheet, full, keys) in enumerate(SUBJECT_SHEETS) for k in keys}


def _key(name):
    return re.sub(r'[^a-z0-9]', '', (name or '').lower())


def subject_sheet(name):
    """(order, sheet title, full name) ya somo — la Halmashauri au jipya."""
    hit = _SUBJECT_LOOKUP.get(_key(name))
    if hit:
        return hit
    title = re.sub(r"[\[\]\*\?/\\:]", ' ', name.upper()).strip()[:31] or 'SUBJECT'
    return (len(SUBJECT_SHEETS), title, name.upper())


def _q(sheet_title):
    """Rejea ya sheet kwenye formula: ='COMP''SCIENCE'!B6 / =KISW!B6."""
    if re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]*', sheet_title):
        return sheet_title
    return "'" + sheet_title.replace("'", "''") + "'"


# ── Titles ──

def _when(joint, months):
    if joint.date:
        return f'{months[joint.date.month - 1]} {joint.date.year}'
    return str(joint.year)


def _title_sw(joint):
    return (f'MUHTASARI WA MATOKEO YA MITIHANI WA {joint.name.upper()}, '
            f'{FORM_SW.get(joint.form, f"KIDATO {joint.form}")} {_when(joint, MONTH_SW)}')


def _title_en(joint):
    region = (joint.region or '').upper()
    prefix = f'{region} REGION - ' if region else ''
    name = joint.name.upper()
    if FORM_EN.get(joint.form, '') not in name:
        name = f'{FORM_EN.get(joint.form, f"FORM {joint.form}")} {name}'
    return f'{prefix}{name}, {_when(joint, MONTH_EN)}'


# ── Template rows ──

class _RowStyles:
    """Mitindo ya mstari wa data na TOTAL kutoka template (kisha inafutwa)."""

    def __init__(self, ws, data_row, total_row, maxcol):
        self.ws, self.maxcol = ws, maxcol
        self.data = [copy.copy(ws.cell(data_row, c)._style) for c in range(1, maxcol + 1)]
        self.total = [copy.copy(ws.cell(total_row, c)._style) for c in range(1, maxcol + 1)]
        self.data_h = ws.row_dimensions[data_row].height
        self.total_h = ws.row_dimensions[total_row].height
        ws.delete_rows(data_row, ws.max_row)

    def row(self, r, total=False):
        for c, s in enumerate(self.total if total else self.data, 1):
            self.ws.cell(r, c)._style = copy.copy(s)
        self.ws.row_dimensions[r].height = self.total_h if total else self.data_h


def _put(ws, r, values):
    for col, v in values.items():
        ws[f'{col}{r}'] = v


# ── DIVISION ──

def _division_sheet(ws, joint, schools):
    first = 9
    styles = _RowStyles(ws, first, first + 1, 44)
    ws['A2'] = f'HALMASHAURI YA WILAYA YA {joint.district.upper()}'
    ws['A3'] = _title_sw(joint)
    last = first + len(schools) - 1
    council = f'{joint.district.upper()} DC'

    for i, s in enumerate(schools):
        r = first + i
        styles.row(r)
        d, a = s['divisions'], s['absent']
        _put(ws, r, {
            'A': i + 1, 'B': council, 'C': s['ward'].upper(), 'D': s['ownership'].upper(),
            'E': s['school'].name.upper() if s['school'] else '',
            'F': f'=IFERROR(I{r}+L{r},"-")', 'G': f'=IFERROR(J{r}+M{r},"-")',
            'H': f'=IF(SUM(F{r}:G{r})=0,"-",SUM(F{r}:G{r}))',
            'I': f'=IFERROR(O{r}+R{r}+U{r}+X{r}+AA{r},"-")',
            'J': f'=IFERROR(P{r}+S{r}+V{r}+Y{r}+AB{r},"-")',
            'K': f'=IF(SUM(I{r}:J{r})=0,"-",SUM(I{r}:J{r}))',
            'L': a['M'], 'M': a['F'], 'N': f'=IF(L{r}="","",SUM(L{r}:M{r}))',
            'O': d['I']['M'], 'P': d['I']['F'], 'Q': f'=IF(O{r}="","",SUM(O{r}:P{r}))',
            'R': d['II']['M'], 'S': d['II']['F'], 'T': f'=IF(R{r}="","",SUM(R{r}:S{r}))',
            'U': d['III']['M'], 'V': d['III']['F'], 'W': f'=IF(U{r}="","",SUM(U{r}:V{r}))',
            'X': d['IV']['M'], 'Y': d['IV']['F'], 'Z': f'=IF(X{r}="","",SUM(X{r}:Y{r}))',
            'AA': d['0']['M'], 'AB': d['0']['F'], 'AC': f'=IF(AA{r}="","",SUM(AA{r}:AB{r}))',
            'AD': f'=IFERROR(O{r}+R{r}+U{r},"")', 'AE': f'=IFERROR(P{r}+S{r}+V{r},"")',
            'AF': f'=IF(AD{r}&AE{r}="","",SUM(AD{r}:AE{r}))',
            'AG': f'=IFERROR(AF{r}/K{r}*100,"-")',
            'AH': f'=IFERROR(AD{r}+X{r},"")', 'AI': f'=IFERROR(AE{r}+Y{r},"")',
            'AJ': f'=IF(AH{r}&AI{r}="","",SUM(AH{r}:AI{r}))',
            'AK': f'=IFERROR(AJ{r}/K{r}*100,"-")',
            'AL': f'=IFERROR(X{r}+AA{r},"")', 'AM': f'=IFERROR(Y{r}+AB{r},"")',
            'AN': f'=IF(AL{r}&AM{r}="","",SUM(AL{r}:AM{r}))',
            'AO': f'=IFERROR(AN{r}/K{r}*100,"-")',
            'AP': f'=IFERROR(Q{r}+T{r}+W{r}+Z{r},"-")',
            'AQ': f'=IFERROR((Q{r}*1+T{r}*2+W{r}*3+Z{r}*4+AC{r}*5)/(Q{r}+T{r}+W{r}+Z{r}+AC{r}),"")',
            'AR': f'=IFERROR(RANK(AQ{r},$AQ${first}:$AQ${last},1),"")',
        })

    if last < first:              # hakuna shule/somo bado — mstari mmoja tupu
        styles.row(first)
        last = first
    t = last + 1
    styles.row(t, total=True)
    ws[f'A{t}'] = 'TOTAL'
    ws.merge_cells(f'A{t}:E{t}')
    for c in range(6, 43):                          # F .. AP
        col = L(c)
        if col in ('AG', 'AK', 'AO'):
            continue
        ws[f'{col}{t}'] = f'=SUM({col}{first}:{col}{last})'
    ws[f'AG{t}'] = f'=IFERROR(AF{t}/K{t}*100,"--")'
    ws[f'AK{t}'] = f'=IFERROR(AJ{t}/K{t}*100,"--")'
    ws[f'AO{t}'] = f'=IFERROR(AN{t}/K{t}*100,"--")'


# ── GRADE ──

def _grade_sheet(ws, joint, schools):
    first = 8
    styles = _RowStyles(ws, first, first + 1, 28)
    ws['B2'] = f'HALMASHAURI YA WILAYA YA {joint.district.upper()}'
    ws['B3'] = _title_sw(joint)
    rows = sorted(schools, key=lambda s: (s['grade']['gpa'] is None, s['grade']['gpa'] or 0))
    last = first + len(rows) - 1
    cols = {'A': ('H', 'I'), 'B': ('K', 'L'), 'C': ('N', 'O'), 'D': ('Q', 'R'), 'F': ('T', 'U')}

    for i, s in enumerate(rows):
        r = first + i
        styles.row(r)
        g = s['grade']['grades']
        vals = {'B': i + 1, 'C': s['ward'].upper(), 'D': s['school'].name.upper() if s['school'] else ''}
        for grade, (fcol, mcol) in cols.items():          # GRADE: F kwanza, kisha M
            vals[fcol], vals[mcol] = g[grade]['F'], g[grade]['M']
        vals.update({
            'E': f'=SUM(H{r},K{r},N{r},Q{r},T{r})', 'F': f'=SUM(I{r},L{r},O{r},R{r},U{r})',
            'G': f'=SUM(E{r}:F{r})',
            'J': f'=SUM(H{r}:I{r})', 'M': f'=SUM(K{r}:L{r})', 'P': f'=SUM(N{r}:O{r})',
            'S': f'=SUM(Q{r}:R{r})', 'V': f'=SUM(T{r}:U{r})',
            'W': f'=H{r}+K{r}+N{r}+Q{r}', 'X': f'=I{r}+L{r}+O{r}+R{r}', 'Y': f'=SUM(W{r}:X{r})',
            'Z': f'=IFERROR(Y{r}/(J{r}+M{r}+P{r}+S{r}+V{r}),"-")',
            'AA': f'=IFERROR((J{r}*1+M{r}*2+P{r}*3+S{r}*4+V{r}*5)/(J{r}+M{r}+P{r}+S{r}+V{r}),"-")',
        })
        _put(ws, r, vals)

    if last < first:              # hakuna shule/somo bado — mstari mmoja tupu
        styles.row(first)
        last = first
    t = last + 1
    styles.row(t, total=True)
    ws[f'C{t}'] = 'TOTAL'
    ws.merge_cells(f'C{t}:D{t}')
    for c in range(5, 26):                          # E .. Y
        col = L(c)
        ws[f'{col}{t}'] = f'=SUM({col}{first}:{col}{last})'
    ws[f'Z{t}'] = f'=IFERROR(Y{t}/(J{t}+M{t}+P{t}+S{t}+V{t}),"-")'
    ws[f'AA{t}'] = f'=IFERROR((J{t}*1+M{t}*2+P{t}*3+S{t}*4+V{t}*5)/(J{t}+M{t}+P{t}+S{t}+V{t}),"-")'


# ── MASOMO ──

GRADE_COLS = {'A': ('H', 'I'), 'B': ('K', 'L'), 'C': ('N', 'O'), 'D': ('Q', 'R'), 'F': ('T', 'U')}


def _subject_sheet(ws, joint, subj):
    first = 9
    styles = _RowStyles(ws, first, first + 1, 28)
    ws['B3'] = f'{joint.district.upper()} DISTRICT COUNCIL'
    ws['B5'] = _title_en(joint)
    ws['B6'] = subj['full_name']
    rows = subj['rows']
    last = first + len(rows) - 1

    for i, row in enumerate(rows):
        r = first + i
        styles.row(r)
        vals = {'B': i + 1, 'C': row['school'].name.upper() if row['school'] else '',
                'D': row['ward'].upper()}
        for grade, (mcol, fcol) in GRADE_COLS.items():    # MASOMO: M kwanza, kisha F
            vals[mcol], vals[fcol] = row['grades'][grade]['M'], row['grades'][grade]['F']
        vals.update({
            'E': f'=SUM(H{r},K{r},N{r},Q{r},T{r})', 'F': f'=SUM(I{r},L{r},O{r},R{r},U{r})',
            'G': f'=SUM(E{r}:F{r})',
            'J': f'=SUM(H{r}:I{r})', 'M': f'=SUM(K{r}:L{r})', 'P': f'=SUM(N{r}:O{r})',
            'S': f'=SUM(Q{r}:R{r})', 'V': f'=SUM(T{r}:U{r})',
            'W': f'=SUM(H{r},K{r},N{r},Q{r})', 'X': f'=SUM(I{r},L{r},O{r},R{r})',
            'Y': f'=SUM(W{r}:X{r})',
            'Z': f'=IFERROR(Y{r}/G{r},"-")',
            'AA': f'=IFERROR((J{r}*1+M{r}*2+P{r}*3+S{r}*4+V{r}*5)/(J{r}+M{r}+P{r}+S{r}+V{r}),"-")',
        })
        _put(ws, r, vals)

    if last < first:              # hakuna shule/somo bado — mstari mmoja tupu
        styles.row(first)
        last = first
    t = last + 1
    styles.row(t, total=True)
    ws[f'C{t}'] = 'TOTAL'
    ws[f'D{t}'] = f'=COUNTA(D{first}:D{last})'
    for c in range(5, 26):                          # E .. Y
        col = L(c)
        ws[f'{col}{t}'] = f'=SUM({col}{first}:{col}{last})'
    ws[f'Z{t}'] = f'=IFERROR(Y{t}/G{t},"-")'
    ws[f'AA{t}'] = f'=IFERROR((J{t}*1+M{t}*2+P{t}*3+S{t}*4+V{t}*5)/(J{t}+M{t}+P{t}+S{t}+V{t}),"-")'
    return t


def _list_sheet(ws, joint, subjects):
    """LIST OF SUBJECTS — kila mstari unasoma TOTAL ya sheet ya somo lake."""
    first = 9
    styles = _RowStyles(ws, first, first + 1, 30)
    ws['C3'] = f'{joint.district.upper()} DISTRICT COUNCIL'
    ws['C4'] = _title_en(joint)
    last = first + len(subjects) - 1

    for i, subj in enumerate(subjects):
        r = first + i
        styles.row(r)
        ref, tot = _q(subj['sheet']), subj['total_row']
        vals = {'C': i + 1, 'D': f'={ref}!$B$6'}
        for c in range(5, 26):                      # E .. Y
            col = L(c)
            vals[col] = f'={ref}!{col}{tot}'
        vals['Z'] = f'=IFERROR(Y{r}/G{r}*100,"-")'
        vals['AA'] = f'={ref}!AA{tot}'
        _put(ws, r, vals)

    if last < first:              # hakuna shule/somo bado — mstari mmoja tupu
        styles.row(first)
        last = first
    t = last + 1
    styles.row(t, total=True)
    ws[f'C{t}'] = 'TOTAL'
    ws.merge_cells(f'C{t}:D{t}')
    for c in range(5, 26):
        col = L(c)
        ws[f'{col}{t}'] = f'=SUM({col}{first}:{col}{last})'
    ws[f'Z{t}'] = f'=IFERROR(Y{t}/G{t}*100,"-")'
    ws[f'AA{t}'] = f'=IFERROR((J{t}*1+M{t}*2+P{t}*3+S{t}*4+V{t}*5)/(J{t}+M{t}+P{t}+S{t}+V{t}),"-")'


# ── Public ──

def _save(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_division_workbook(joint, ownership=None) -> bytes:
    """DIVISION + GRADE (format ya "FORM ONE DIVISION PERFORMANCE ANALYSIS")."""
    data = analyse_joint_exam(joint, ownership=ownership)
    wb = load_workbook(TEMPLATE_DIR / 'joint_division.xlsx')
    _division_sheet(wb['DIVISION'], joint, data['schools'])
    _grade_sheet(wb['GRADE'], joint, data['schools'])
    return _save(wb)


def build_subjects_workbook(joint, ownership=None) -> bytes:
    """LIST OF SUBJECTS + sheet ya kila somo (format ya "SUBJECT PERFORMANCE ANALYSIS")."""
    data = analyse_joint_exam(joint, ownership=ownership)
    wb = load_workbook(TEMPLATE_DIR / 'joint_subjects.xlsx')
    template = wb['SUBJECT_TEMPLATE']

    subjects = []
    used = set()
    for s in data['subjects']:
        order, sheet, full = subject_sheet(s['name'])
        base, n = sheet, 2
        while sheet in used:
            sheet = f'{base[:28]} {n}'
            n += 1
        used.add(sheet)
        subjects.append({**s, 'order': order, 'sheet': sheet, 'full_name': full})

    # Sheets kwa mpangilio wa faili la Halmashauri (H'MAADILI, PHYS, EDK, ...)
    for subj in sorted(subjects, key=lambda x: (x['order'], x['sheet'])):
        ws = wb.copy_worksheet(template)
        ws.title = subj['sheet']
        ws.sheet_view.zoomScale = template.sheet_view.zoomScale
        subj['total_row'] = _subject_sheet(ws, joint, subj)
    wb.remove(template)

    # LIST: kwa GPA (bora kwanza), kama kwenye faili la Halmashauri
    ranked = sorted(subjects, key=lambda x: (x['gpa'] is None, x['gpa'] or 0, x['full_name']))
    _list_sheet(wb['LIST OF SUBJECTS'], joint, ranked)
    return _save(wb)


def workbook_filename(joint, kind, ownership=None):
    form = FORM_EN.get(joint.form, f'FORM {joint.form}')
    own = OWNERSHIP_LABEL.get(ownership, 'ALL')
    if kind == 'division':
        name = f'{form} DIVISION PERFORMANCE ANALYSIS {joint.year} - {own}.xlsx'
    else:
        name = f'SUBJECT PERFORMANCE ANALYSIS_{form} - {own} - {joint.name.upper()} {joint.year}.xlsx'
    return re.sub(r'[\\/:*?"<>|]', '-', name)
